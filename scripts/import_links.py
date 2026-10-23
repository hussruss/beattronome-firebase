#!/usr/bin/env python3
"""Load an Excel's links as pending tasks for the backoffice "Ligas" page
and its sendLinkBatch push. One-off: delete with the module.

Setup (once):
  python3 -m venv /tmp/seedenv
  /tmp/seedenv/bin/pip install firebase-admin openpyxl

Usage:
  GOOGLE_APPLICATION_CREDENTIALS=~/Downloads/<beattronome-sa>.json \\
    /tmp/seedenv/bin/python3 scripts/import_links.py <file.xlsx> \\
    --sheet "Cruce (No me siguen)" --owner hussrusse@gmail.com

  --dry-run      print what would be imported, write nothing
  --days N       days to finish, from today (default 75); only changes an
                 existing deadline when passed
  --per-day N    notifications per day (default 7)

Every cell with a hyperlink or an http(s) URL (typed or a formula's cached
result) is a link; the row's first other text is its label. Re-running only
adds new URLs: links already imported (done or not) are left as they are.
The project is the service account's.
"""
import argparse
import hashlib
import random
import re
import sys
from datetime import datetime, timedelta, timezone

import openpyxl

URL = re.compile(r"^https?://\S+$", re.IGNORECASE)
# America/Mexico_City has no DST since 2022.
MEXICO = timezone(timedelta(hours=-6))


def read_links(path, only_sheet):
    # data_only: formulas come back as their cached result.
    book = openpyxl.load_workbook(path, data_only=True)
    if only_sheet and only_sheet not in book.sheetnames:
        sys.exit(f"Sheet not found: {only_sheet} (sheets: {book.sheetnames})")
    sheets = [book[only_sheet]] if only_sheet else book.worksheets
    seen = {}
    for sheet in sheets:
        for row in sheet.iter_rows():
            urls, label = [], None
            for cell in row:
                text = str(cell.value).strip() if cell.value is not None else ""
                link = cell.hyperlink.target.strip() if cell.hyperlink and cell.hyperlink.target else None
                if link and URL.match(link):
                    urls.append(link)
                    if not label and text and not URL.match(text):
                        label = text
                elif URL.match(text):
                    urls.append(text)
                elif not label and text:
                    label = text
            for url in urls:
                seen.setdefault(url, {"url": url, "label": label, "sheet": sheet.title, "row": row[0].row})
    return list(seen.values())


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("file")
    p.add_argument("--sheet")
    p.add_argument("--owner")
    p.add_argument("--days", type=int)
    p.add_argument("--per-day", type=int, default=7)
    p.add_argument("--dry-run", action="store_true")
    a = p.parse_args()

    links = read_links(a.file, a.sheet)
    print(f"{len(links)} distinct links in {a.file}")
    for l in links[:5]:
        print(f"  {l['sheet']}!{l['row']}  {l['label'] or ''}  {l['url']}")
    if a.dry_run:
        return
    if not a.owner:
        sys.exit("--owner is required")

    import firebase_admin
    from firebase_admin import auth, firestore
    from google.api_core.exceptions import AlreadyExists

    app = firebase_admin.initialize_app()
    print(f"Project: {app.project_id}")
    owner_uid = auth.get_user_by_email(a.owner).uid if "@" in a.owner else a.owner

    db = firestore.client()
    col = db.collection("link_tasks")
    added = existing = 0
    for i, l in enumerate(links, 1):
        doc_id = hashlib.sha1(l["url"].encode()).hexdigest()[:20]
        data = {
            "url": l["url"],
            "sheet": l["sheet"],
            "row": l["row"],
            "status": "pending",
            "rand": random.random(),
            "unsentRand": random.random(),
            "createdAt": firestore.SERVER_TIMESTAMP,
        }
        if l["label"]:
            data["label"] = l["label"]
        try:
            col.document(doc_id).create(data)
            added += 1
        except AlreadyExists:
            existing += 1
        if i % 250 == 0:
            print(f"  {i}/{len(links)}")

    cfg_ref = db.collection("link_tasks_config").document("state")
    cfg = cfg_ref.get().to_dict() or {}
    if a.days is not None or not cfg.get("deadline"):
        deadline = (datetime.now(MEXICO) + timedelta(days=a.days or 75)).strftime("%Y-%m-%d")
    else:
        deadline = cfg["deadline"]
    cfg_ref.set({"ownerUid": owner_uid, "deadline": deadline, "perDay": a.per_day, "enabled": True}, merge=True)
    print(f"Added {added}, already there {existing}. Deadline {deadline}, "
          f"{a.per_day} notifications a day to {a.owner}.")


if __name__ == "__main__":
    main()
